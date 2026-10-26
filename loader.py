# -*- coding: utf-8 -*-
"""
loader.py（v3.1）
职责：读取 PDF / PPTX 并切分文本片段（单一职责，不涉及向量化和问答）。
- PDF：PyPDF2 逐页提取文本；无文本页（扫描件、图片页）自动走 PyMuPDF + RapidOCR 兜底；
- PPTX：python-pptx 逐张幻灯片提取文本（含表格单元格文本）；
- 清洗文本：合并段内换行、压缩多余空白；
- 按「句子边界优先」切分：每段约 500 字，段间保留约 50 字的句子级重叠；
- 单句超过 500 字时，退化为字符滑动窗口硬切兜底。

v3.1 新增：load_and_split_pptx、load_and_split_file（按后缀分发）、无文本页 OCR 兜底。
"""
import io
import os
import re
from typing import IO, Union

from PyPDF2 import PdfReader

# 切分参数：每段目标长度 / 段间重叠长度（单位：字符）
CHUNK_SIZE = 500
CHUNK_OVERLAP = 50

# 单页清洗后文本少于该字符数，就判定为"无文本页"并触发 OCR 兜底
MIN_TEXT_LEN = 10
# OCR 前把页面渲染成图片的缩放倍数（2 倍清晰度足够，体积也可控）
OCR_ZOOM = 2.0
# 目前支持的文件后缀（小写）
SUPPORTED_EXTS = ("pdf", "pptx")

# OCR 引擎懒加载缓存：只有真的遇到无文本页时才会初始化（首次会下载识别模型）
_ocr_engine = None
_ocr_error = ""


def _clean_text(text: str) -> str:
    """清洗单页文本：合并段内换行、去掉多余空格。"""
    if not text:
        return ""
    # 统一换行符
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    # 多个连续换行视为段落分隔，压缩为一个
    text = re.sub(r"\n{2,}", "\n", text)
    # 中文 PDF 的段内换行多为排版产生，直接去掉以免切断句子
    text = text.replace("\n", "")
    # 压缩连续空格/制表符
    text = re.sub(r"[ \t]+", " ", text)
    return text.strip()


def _split_sentences(text: str) -> list[str]:
    """按中英文句末标点切句，标点保留在句尾。"""
    sentences = re.split(r"(?<=[。！？；!?;])", text)
    return [s for s in sentences if s.strip()]


def _hard_split(sentence: str) -> list[str]:
    """兜底：对超长单句使用 500 字 / 50 字重叠的字符滑动窗口硬切。"""
    chunks = []
    step = CHUNK_SIZE - CHUNK_OVERLAP
    for start in range(0, len(sentence), step):
        chunk = sentence[start:start + CHUNK_SIZE]
        if chunk.strip():
            chunks.append(chunk)
        if start + CHUNK_SIZE >= len(sentence):
            break
    return chunks


def _split_page_text(text: str) -> list[str]:
    """把单页清洗后的文本切成约 500 字的片段，段间保留约 50 字句子级重叠。"""
    sentences = _split_sentences(text)
    chunks = []
    current = ""

    for sentence in sentences:
        # 单句本身超长：先结算当前段落，再对该句硬切
        if len(sentence) > CHUNK_SIZE:
            if current:
                chunks.append(current)
                current = ""
            chunks.extend(_hard_split(sentence))
            continue

        if len(current) + len(sentence) <= CHUNK_SIZE:
            current += sentence
        else:
            chunks.append(current)
            # 从刚完成的段落尾部取若干句子，凑出约 50 字重叠，作为下一段开头
            overlap = ""
            for s in reversed(_split_sentences(current)):
                if len(overlap) + len(s) > CHUNK_OVERLAP:
                    break
                overlap = s + overlap
            current = overlap + sentence

    if current.strip():
        chunks.append(current)
    return chunks


def _read_file_bytes(file: Union[str, IO[bytes]]) -> bytes:
    """
    统一把入参读成 bytes：既支持文件路径，也支持 Streamlit 的 UploadedFile。
    读成 bytes 后可以反复使用（PyPDF2 与 PyMuPDF 各解析一次），避免流指针位置问题。
    """
    if isinstance(file, (str, bytes, os.PathLike)):
        path = os.fspath(file)
        with open(path, "rb") as f:
            return f.read()
    try:
        file.seek(0)
    except Exception:
        pass  # 不支持 seek 的流对象直接读即可
    data = file.read()
    if isinstance(data, str):
        data = data.encode("utf-8")
    return data


def _get_ocr():
    """
    懒加载 RapidOCR 引擎（只有真的遇到无文本页时才会调用）。
    首次调用会自动下载识别模型（约几十 MB），失败时抛 RuntimeError 并缓存原因，
    避免对每一页重复尝试、拖慢入库。
    """
    global _ocr_engine, _ocr_error
    if _ocr_engine is not None:
        return _ocr_engine
    if _ocr_error:
        raise RuntimeError(_ocr_error)
    try:
        from rapidocr_onnxruntime import RapidOCR

        _ocr_engine = RapidOCR()
    except Exception as e:  # 模型未下载、依赖缺失等情况都走这里
        _ocr_error = f"OCR 引擎初始化失败：{e}"
        raise RuntimeError(_ocr_error) from e
    return _ocr_engine


def _ocr_image(image_bytes: bytes) -> str:
    """对单张图片做文字识别，返回识别出的纯文本（失败返回空串）。"""
    engine = _get_ocr()
    result, _ = engine(image_bytes)
    if not result:
        return ""
    # RapidOCR 返回结构：[[框坐标, 文本, 置信度], ...]
    return "\n".join(str(item[1]) for item in result if len(item) > 1)


def _render_pdf_page_png(data: bytes, page_no: int, zoom: float = OCR_ZOOM) -> bytes:
    """用 PyMuPDF 把 PDF 指定页渲染成 PNG 图片字节（zoom 为缩放倍数）。"""
    try:  # 新版 PyMuPDF 推荐 import pymupdf，旧版只有 fitz
        import pymupdf as fitz
    except ImportError:
        import fitz
    doc = fitz.open(stream=data, filetype="pdf")
    try:
        page = doc[page_no - 1]  # PyMuPDF 页码从 0 开始
        pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom))
        return pix.tobytes("png")
    finally:
        doc.close()


def load_and_split_pdf(
    file: Union[str, IO[bytes]], ocr_report: dict = None
) -> list[dict]:
    """
    读取 PDF 并切分文本；无文本页（扫描件、图片页）自动走 OCR 兜底。
    :param file: PDF 文件路径，或类文件对象（如 Streamlit 的 UploadedFile）
    :param ocr_report: 可选，传入一个空 dict，函数会把 OCR 情况写进去：
                       {"ocr_pages": [走过OCR的页码], "ocr_failed": [OCR仍失败的页码],
                        "ocr_error": "引擎不可用的原因"}
    :return: [{"text": 片段文本, "page": 页码(从1开始)}, ...]
    """
    reader = PdfReader(file)
    results = []
    pending_ocr = []  # 无文本页：先记下页码，稍后统一 OCR

    def _append(page_no: int, text: str):
        """把一页文本切分后追加进结果（空文本直接跳过）。"""
        if not text:
            return
        for chunk in _split_page_text(text):
            results.append({"text": chunk, "page": page_no})

    for page_no, page in enumerate(reader.pages, start=1):
        cleaned = _clean_text(page.extract_text() or "")
        if len(cleaned) < MIN_TEXT_LEN:
            # 记录页码与"原本那点文本"：OCR 失败时还能退回用它，不把已有内容弄丢
            pending_ocr.append((page_no, cleaned))
            continue
        _append(page_no, cleaned)

    ocr_pages, ocr_failed = [], []
    if pending_ocr:
        # 只有在真的存在无文本页时，才把原始数据读成 bytes 交给 PyMuPDF 渲染，
        # 避免正常 PDF 被多余地整份读进内存。
        try:
            data = _read_file_bytes(file)
        except Exception as e:
            if ocr_report is not None:
                ocr_report["ocr_error"] = f"无法读取原始文件，OCR 跳过：{e}"
            data = None

        if data:
            for page_no, fallback in pending_ocr:
                try:
                    png = _render_pdf_page_png(data, page_no)
                    text = _clean_text(_ocr_image(png))
                except Exception as e:
                    # 单页 OCR 失败绝不能中断整个文件的入库
                    if ocr_report is not None:
                        ocr_report["ocr_error"] = str(e)
                    text = ""
                if len(text) >= MIN_TEXT_LEN:
                    ocr_pages.append(page_no)
                    _append(page_no, text)
                elif fallback:
                    # OCR 没帮上忙，但原本还提取到一点文本：保留它，不算跳过
                    _append(page_no, fallback)
                else:
                    # 真的一个字都没有，才算这一页被跳过
                    ocr_failed.append(page_no)
        else:
            # 读不到原始数据：退回原本提取到的短文本，只有完全无文本的页才算跳过
            for page_no, fallback in pending_ocr:
                if fallback:
                    _append(page_no, fallback)
                else:
                    ocr_failed.append(page_no)

    if ocr_report is not None:
        ocr_report["ocr_pages"] = ocr_pages
        ocr_report["ocr_failed"] = ocr_failed

    # OCR 的页是后追加的，按页码重新排一次，保证输出顺序与原文一致
    results.sort(key=lambda c: c["page"])
    return results


def _iter_shape_text(shape) -> list[str]:
    """
    递归取一个形状里的文本：优先表格单元格，其次文本框，组合形状继续下钻。
    返回文本列表（每张幻灯片内部再用换行连接）。
    """
    texts = []
    # 表格：逐单元格取文本
    if getattr(shape, "has_table", False):
        try:
            for row in shape.table.rows:
                for cell in row.cells:
                    if cell.text and cell.text.strip():
                        texts.append(cell.text.strip())
            return texts
        except Exception:
            return texts

    # 普通文本框（标题、正文、项目符号等）
    if getattr(shape, "has_text_frame", False):
        text = shape.text_frame.text
        if text and text.strip():
            texts.append(text.strip())
        return texts

    # 组合形状：递归子形状
    sub_shapes = getattr(shape, "shapes", None)
    if sub_shapes is not None:
        try:
            for sub in sub_shapes:
                texts.extend(_iter_shape_text(sub))
        except Exception:
            pass
    return texts


def load_and_split_pptx(file: Union[str, IO[bytes]]) -> list[dict]:
    """
    读取 PPTX 并切分文本：逐张幻灯片提取所有形状的文本（含表格单元格），
    同一张幻灯片内的文本用换行连接，然后复用清洗与切分逻辑。
    :param file: PPTX 文件路径，或类文件对象（如 Streamlit 的 UploadedFile）
    :return: [{"text": 片段文本, "page": 幻灯片序号(从1开始)}, ...]
             与 load_and_split_pdf 完全一致，便于入库与引用逻辑复用
    """
    from pptx import Presentation

    data = _read_file_bytes(file)
    prs = Presentation(io.BytesIO(data))
    results = []
    for slide_no, slide in enumerate(prs.slides, start=1):
        parts = []
        for shape in slide.shapes:
            parts.extend(_iter_shape_text(shape))
        cleaned = _clean_text("\n".join(parts))
        if not cleaned:
            continue  # 跳过空白幻灯片
        for chunk in _split_page_text(cleaned):
            results.append({"text": chunk, "page": slide_no})
    return results


def load_and_split_file(
    file: Union[str, IO[bytes]], filename: str, ocr_report: dict = None
) -> list[dict]:
    """
    按文件后缀分发解析器（PDF 走 PDF 解析，PPTX 走 PPTX 解析）。
    :param filename: 原始文件名，仅用于判断后缀
    :param ocr_report: 仅 PDF 使用，见 load_and_split_pdf 的说明
    :raises ValueError: 遇到不支持的后缀
    """
    ext = os.path.splitext(filename or "")[1].lower().lstrip(".")
    if ext == "pdf":
        return load_and_split_pdf(file, ocr_report=ocr_report)
    if ext == "pptx":
        return load_and_split_pptx(file)
    raise ValueError(f"不支持的文件类型：.{ext or '未知'}（目前支持 pdf / pptx）")
