# -*- coding: utf-8 -*-
"""
图片型 PDF / 扫描件 OCR 兜底的单元测试（v3.1 新增）

测试原则：
1. 用 PyMuPDF 程序化生成"只有图片、没有文本层"的 PDF，不依赖外部素材；
2. 默认把识别函数 _ocr_image 换成假实现，因此**不会真的下载 OCR 模型**，
   用例稳定且秒级完成；
3. 真实 OCR 引擎的用例默认跳过（需设置环境变量 RUN_REAL_OCR=1 才执行），
   引擎不可用时 pytest.skip 会给出明确原因。
"""
import os
from unittest.mock import patch

import pytest

import loader


def _open_fitz():
    """兼容新旧两种 PyMuPDF 导入方式（新版推荐 pymupdf，旧版只有 fitz）。"""
    try:
        import pymupdf as fitz
    except ImportError:
        import fitz
    return fitz


@pytest.fixture
def image_pdf(tmp_path):
    """
    生成一个"单页纯图片"PDF：页面里只插入一张图片，没有任何文本层，
    因此 PyPDF2 的 extract_text() 会返回空串，从而触发 OCR 兜底分支。
    """
    fitz = _open_fitz()
    doc = fitz.open()
    page = doc.new_page(width=600, height=400)
    # 把当前空白页渲染成 PNG，再当作图片插回去（得到一页"扫描件式的图片页"）
    png = page.get_pixmap().tobytes("png")
    page.insert_image(fitz.Rect(50, 50, 550, 350), stream=png)
    path = str(tmp_path / "scan.pdf")
    doc.save(path)
    doc.close()
    return path


# ---------------- 走 OCR 分支 ----------------

def test_图片页会走OCR分支返回文本(image_pdf):
    """无文本页应触发 OCR，并把识别出的文本按原页码返回。"""
    fake_text = "这是OCR识别出来的扫描件正文，长度足够触发切分。"
    with patch("loader._ocr_image", return_value=fake_text) as mock_ocr:
        chunks = loader.load_and_split_pdf(image_pdf)

    assert mock_ocr.called, "无文本页应当调用 OCR"
    assert chunks, "OCR 识别出文本后应有片段返回"
    assert all(c["page"] == 1 for c in chunks)
    assert "OCR识别" in "".join(c["text"] for c in chunks)


def test_OCR结果也会走清洗与切分(image_pdf):
    """OCR 出来的文本必须同样经过 _clean_text 与 _split_page_text 处理。"""
    fake_text = "\n这是   带排版换行与多余空格的识别结果。\n第二句内容在这里。\n"
    with patch("loader._ocr_image", return_value=fake_text):
        chunks = loader.load_and_split_pdf(image_pdf)

    for c in chunks:
        assert "\n" not in c["text"], "清洗后不应残留换行"
        assert "    " not in c["text"], "清洗后不应残留连续空格"


def test_OCR识别不出内容时记录跳过的页码(image_pdf):
    """OCR 返回空文本时：结果为空，且被跳过的页码写进 ocr_report。"""
    report = {}
    with patch("loader._ocr_image", return_value=""):
        chunks = loader.load_and_split_pdf(image_pdf, ocr_report=report)

    assert chunks == []
    assert report.get("ocr_failed") == [1]


def test_OCR报错不会中断入库(image_pdf):
    """OCR 引擎不可用时（如模型未下载）不能抛异常，必须记录原因并跳过该页。"""
    report = {}
    with patch("loader._ocr_image", side_effect=RuntimeError("模型未下载")):
        chunks = loader.load_and_split_pdf(image_pdf, ocr_report=report)

    assert chunks == []
    assert report.get("ocr_failed") == [1]
    assert "模型未下载" in report.get("ocr_error", "")


def test_单页OCR失败不影响其他页(tmp_path):
    """多页 PDF 中一页 OCR 失败，其他页的正常文本仍要保留。"""
    fitz = _open_fitz()
    doc = fitz.open()
    p1 = doc.new_page()
    p1.insert_text((72, 120), "第一页是正常文本层的内容，长度足够。")
    p2 = doc.new_page()
    png = p2.get_pixmap().tobytes("png")
    p2.insert_image(fitz.Rect(50, 50, 500, 350), stream=png)  # 第二页只有图片
    path = str(tmp_path / "mixed.pdf")
    doc.save(path)
    doc.close()

    report = {}
    with patch("loader._ocr_image", side_effect=RuntimeError("引擎不可用")):
        chunks = loader.load_and_split_pdf(path, ocr_report=report)

    assert {c["page"] for c in chunks} == {1}, "只有第一页有可用文本"
    assert report.get("ocr_failed") == [2]


# ---------------- 正常页不触发 OCR ----------------

def test_有正常文本的页不会触发OCR(tmp_path):
    """关键约束：有文本的页一律不走 OCR，避免拖慢普通 PDF。"""
    fitz = _open_fitz()
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 120), "这一页有正常的文本层，内容长度足够，不需要 OCR。")
    path = str(tmp_path / "text.pdf")
    doc.save(path)
    doc.close()

    with patch("loader._ocr_image") as mock_ocr:
        chunks = loader.load_and_split_pdf(path)

    assert chunks
    assert not mock_ocr.called, "有文本的页不应调用 OCR"


def test_PPTX解析不依赖OCR(tmp_path):
    """PPTX 走独立解析路径，不涉及 OCR。"""
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "PPT标题内容足够长以通过清洗"
    path = str(tmp_path / "d.pptx")
    prs.save(path)

    with patch("loader._ocr_image") as mock_ocr:
        chunks = loader.load_and_split_pptx(path)

    assert chunks
    assert not mock_ocr.called


# ---------------- 真实 OCR 引擎（默认跳过） ----------------

def test_真实OCR引擎可用时图片页能正常处理(image_pdf):
    """
    真实 OCR 冒烟用例：默认跳过（首次运行需下载识别模型，耗时较长）。
    设置环境变量 RUN_REAL_OCR=1 可启用；引擎不可用时也会跳过并说明原因。
    """
    if os.environ.get("RUN_REAL_OCR") != "1":
        pytest.skip(
            "跳过真实 OCR 用例：默认不下载识别模型。"
            "设置环境变量 RUN_REAL_OCR=1 后可启用（首次会自动下载模型，耗时较长）。"
        )
    try:
        loader._get_ocr()
    except Exception as e:
        pytest.skip(f"跳过：OCR 引擎不可用（{e}）。原因通常是识别模型未下载或依赖未安装。")

    report = {}
    chunks = loader.load_and_split_pdf(image_pdf, ocr_report=report)

    # 空白图片可能识别不出文字，因此只断言：不抛异常、返回列表、报告结构完整
    assert isinstance(chunks, list)
    assert "ocr_failed" in report
