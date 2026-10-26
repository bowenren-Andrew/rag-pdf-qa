# -*- coding: utf-8 -*-
"""
PPTX 解析的单元测试（v3.1 新增）

测试原则：
1. 用 python-pptx 在测试里**程序化生成** pptx 文件，不依赖外部素材；
2. 只测 loader 的解析与切分，不 import vectorstore，因此不会触发向量模型下载；
3. 不涉及任何网络请求与真实 API 调用。
"""
import pytest
from pptx import Presentation
from pptx.util import Inches

from loader import load_and_split_file, load_and_split_pptx

# 第一张幻灯片的标题与正文、第二张的标题与表格内容（断言时用）
TITLE_1 = "机器学习基础知识"
BODY_1 = "监督学习需要标注数据。无监督学习不需要标注数据。"
TITLE_2 = "算法效果对比"
CELL_TEXT = "决策树算法准确率"


def _build_sample_pptx(path) -> str:
    """程序化生成一个 2 页 pptx：第 1 页标题+正文，第 2 页标题+2×2 表格。"""
    prs = Presentation()

    # 版式 1 = "标题和内容"
    slide1 = prs.slides.add_slide(prs.slide_layouts[1])
    slide1.shapes.title.text = TITLE_1
    slide1.placeholders[1].text_frame.text = BODY_1

    # 版式 5 = "仅标题"，再手工插入一个表格
    slide2 = prs.slides.add_slide(prs.slide_layouts[5])
    slide2.shapes.title.text = TITLE_2
    shape = slide2.shapes.add_table(
        rows=2, cols=2, left=Inches(1), top=Inches(2), width=Inches(6), height=Inches(1.5)
    )
    table = shape.table
    table.cell(0, 0).text = "算法"
    table.cell(0, 1).text = "准确率"
    table.cell(1, 0).text = CELL_TEXT
    table.cell(1, 1).text = "92%"

    prs.save(str(path))
    return str(path)


@pytest.fixture
def sample_pptx(tmp_path):
    """生成一个临时 pptx 文件并返回其路径。"""
    return _build_sample_pptx(tmp_path / "demo.pptx")


# ---------------- load_and_split_pptx ----------------

def test_解析PPTX得到两张幻灯片的页码(sample_pptx):
    """两页 PPT 应被解析成 page=1 与 page=2 的片段。"""
    chunks = load_and_split_pptx(sample_pptx)

    assert chunks, "PPTX 应至少解析出一个片段"
    assert {c["page"] for c in chunks} == {1, 2}


def test_解析结果包含标题与正文文本(sample_pptx):
    """标题、正文都要能被提取到（各形状的文本合并后清洗切分）。"""
    chunks = load_and_split_pptx(sample_pptx)
    all_text = "".join(c["text"] for c in chunks)

    assert TITLE_1 in all_text
    assert "监督学习" in all_text
    assert TITLE_2 in all_text


def test_表格单元格文本也能被提取(sample_pptx):
    """第二页是表格，单元格里的文字同样要进入片段。"""
    chunks = load_and_split_pptx(sample_pptx)
    all_text = "".join(c["text"] for c in chunks)

    assert CELL_TEXT in all_text
    assert "92%" in all_text


def test_PPTX返回结构与PDF完全一致(sample_pptx):
    """返回结构必须兼容入库逻辑：每条都是 {'text':..., 'page':...}。"""
    chunks = load_and_split_pptx(sample_pptx)

    for c in chunks:
        assert set(c.keys()) == {"text", "page"}
        assert isinstance(c["page"], int)
        assert c["text"].strip()


def test_页码与顺序按幻灯片递增(sample_pptx):
    """片段的页码应当单调递增，便于来源展示「第 X 张」。"""
    chunks = load_and_split_pptx(sample_pptx)
    pages = [c["page"] for c in chunks]

    assert pages == sorted(pages)
    assert pages[0] == 1


# ---------------- load_and_split_file 分发 ----------------

def test_按后缀分发到PPTX解析器(sample_pptx):
    """文件名以 .pptx 结尾时应走 PPTX 解析，结果与直接调用一致。"""
    by_file = load_and_split_file(sample_pptx, "demo.pptx")
    direct = load_and_split_pptx(sample_pptx)

    assert by_file == direct


def test_不支持的后缀给出明确错误(sample_pptx):
    """遇到不支持的类型应抛 ValueError，且错误信息里带上后缀名。"""
    with pytest.raises(ValueError, match="不支持的文件类型"):
        load_and_split_file(sample_pptx, "demo.docx")


def test_后缀判断不区分大小写(sample_pptx):
    """大写后缀 .PPTX 也应正确识别。"""
    chunks = load_and_split_file(sample_pptx, "DEMO.PPTX")

    assert chunks, "大写后缀应同样被识别为 PPTX"
