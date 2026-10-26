# -*- coding: utf-8 -*-
"""
强制全文 OCR（force_ocr）的单元测试（v3.2 新增）

背景：LaTeX 导出的公式类 PDF 能提取到文字（整页远超 10 字），但公式符号没有
文本映射、提取后公式位置为空白，v3.1 的"无文本页才 OCR"规则不会触发。
因此提供 force_ocr 开关：忽略文本层，逐页渲染成图片走 RapidOCR。

测试原则（与 v3.1 的 test_ocr_fallback.py 一致）：
1. 用 PyMuPDF 程序化生成 PDF，不依赖外部素材；
2. 默认把 _ocr_image 换成假实现，因此**不会真的加载 OCR 模型**，秒级完成；
3. 只 gating 一个真实 OCR 冒烟用例（RUN_REAL_OCR=1 才执行）。
"""
import os
from unittest.mock import patch

import pytest

import loader


def _open_fitz():
    """兼容新旧两种 PyMuPDF 导入方式（新版推荐 pymupdf，旧版只有 fitz）。"""
    try:
        import pymupdf as fitz
    except ImportError:
        import fitz
    return fitz


@pytest.fixture
def text_pdf(tmp_path):
    """
    生成一个"2 页文本层 PDF"：每页都有充足文字（远超 MIN_TEXT_LEN），
    因此默认（force_ocr=False）时**不会**触发 OCR——用来验证强制模式确实改变了行为。
    """
    fitz = _open_fitz()
    doc = fitz.open()
    p1 = doc.new_page()
    p1.insert_text((72, 120), "Page one has a normal text layer, long enough.")
    p2 = doc.new_page()
    p2.insert_text((72, 120), "Page two also has a normal text layer, long enough.")
    path = str(tmp_path / "text.pdf")
    doc.save(path)
    doc.close()
    return path


@pytest.fixture
def image_pdf(tmp_path):
    """生成一个"单页纯图片"PDF：没有文本层，OCR 失败时连回退文本都没有。"""
    fitz = _open_fitz()
    doc = fitz.open()
    page = doc.new_page(width=600, height=400)
    png = page.get_pixmap().tobytes("png")
    page.insert_image(fitz.Rect(50, 50, 550, 350), stream=png)
    path = str(tmp_path / "scan.pdf")
    doc.save(path)
    doc.close()
    return path


# ---------------- 强制 OCR 的基本行为 ----------------

def test_强制OCR时每页都会调用OCR(text_pdf):
    """force_ocr=True 时忽略文本层，每一页都应送进 OCR（这是本次要解决的核心问题）。"""
    fake_text = "这是OCR识别出来的正文内容，长度足够触发切分逻辑。"
    with patch("loader._ocr_image", return_value=fake_text) as mock_ocr:
        report = {}
        chunks = loader.load_and_split_pdf(text_pdf, ocr_report=report, force_ocr=True)

    assert mock_ocr.call_count == 2, "两页都应调用 OCR"
    assert report.get("ocr_pages") == [1, 2], "两页都应记入 ocr_pages"
    assert {c["page"] for c in chunks} == {1, 2}, "两页都应有片段产出"
    assert "OCR识别" in "".join(c["text"] for c in chunks), "片段内容应来自 OCR 结果"


def test_强制OCR忽略文本层(text_pdf):
    """强制模式下优先采用 OCR 结果，而不是原本提取到的文本层内容。"""
    with patch("loader._ocr_image", return_value="OCR替换掉的内容。") as _:
        chunks = loader.load_and_split_pdf(text_pdf, force_ocr=True)

    joined = "".join(c["text"] for c in chunks)
    assert "Page one has a normal text layer" not in joined, "不应再出现文本层原文"
    assert "OCR替换掉的内容" in joined


def test_强制OCR的结果仍按页码排序(text_pdf):
    """OCR 页是后追加的，返回前必须按页码重新排序。"""
    with patch("loader._ocr_image", return_value="一段足够长的识别文本，用于切分校验。"):
        chunks = loader.load_and_split_pdf(text_pdf, force_ocr=True)

    pages = [c["page"] for c in chunks]
    assert pages == sorted(pages), "片段必须按页码升序"


def test_强制OCR结果也会走清洗与切分(text_pdf):
    """OCR 出来的文本必须同样经过 _clean_text 与 _split_page_text 处理。"""
    fake_text = "\n带排版换行与   多余空格的识别结果。\n第二句内容在这里。\n"
    with patch("loader._ocr_image", return_value=fake_text):
        chunks = loader.load_and_split_pdf(text_pdf, force_ocr=True)

    assert chunks
    for c in chunks:
        assert "\n" not in c["text"], "清洗后不应残留换行"
        assert "    " not in c["text"], "清洗后不应残留连续空格"


def test_ocr_report会标记强制模式(text_pdf):
    """report 里要能区分 auto / force，供界面给出不同的提示文案。"""
    with patch("loader._ocr_image", return_value="识别文本足够长以通过阈值校验。"):
        report = {}
        loader.load_and_split_pdf(text_pdf, ocr_report=report, force_ocr=True)
        assert report.get("ocr_mode") == "force"

        report2 = {}
        loader.load_and_split_pdf(text_pdf, ocr_report=report2)
        assert report2.get("ocr_mode") == "auto"


# ---------------- 强制 OCR 的失败回退 ----------------

def test_强制OCR单页失败时回退提取文本(text_pdf):
    """单页 OCR 失败不能中断整文件：应回退到该页提取到的文本，且不算跳过。"""
    report = {}
    with patch("loader._ocr_image", side_effect=RuntimeError("引擎不可用")):
        chunks = loader.load_and_split_pdf(text_pdf, ocr_report=report, force_ocr=True)

    joined = "".join(c["text"] for c in chunks)
    assert "Page one has a normal text layer" in joined, "应回退到文本层内容"
    assert {c["page"] for c in chunks} == {1, 2}
    assert report.get("ocr_failed") == [], "有文本可回退的页不算失败"
    assert report.get("ocr_pages") == [], "失败的页不应记入成功列表"
    assert "引擎不可用" in report.get("ocr_error", "")


def test_强制OCR失败且无文本可回退时记录页码(image_pdf):
    """既没有 OCR 结果、也没有文本层时，才把页码记入 ocr_failed。"""
    report = {}
    with patch("loader._ocr_image", side_effect=RuntimeError("模型未下载")):
        chunks = loader.load_and_split_pdf(image_pdf, ocr_report=report, force_ocr=True)

    assert chunks == []
    assert report.get("ocr_failed") == [1], "无内容可回退的页才算跳过"


def test_强制OCR部分页失败不影响其他页(tmp_path):
    """多页中只有一页 OCR 失败，其他页的正常识别结果仍要保留。"""
    fitz = _open_fitz()
    doc = fitz.open()
    for _ in range(3):
        page = doc.new_page()
        page.insert_text((72, 120), "Some text layer content long enough here.")
    path = str(tmp_path / "three.pdf")
    doc.save(path)
    doc.close()

    def _fake_ocr(png_bytes):
        # 让第二页 OCR 失败，模拟单页识别异常
        if _fake_ocr.calls == 1:
            _fake_ocr.calls += 1
            raise RuntimeError("第二页识别失败")
        _fake_ocr.calls += 1
        return "这一页OCR识别成功，内容长度足够用于切分。"

    _fake_ocr.calls = 0
    report = {}
    with patch("loader._ocr_image", side_effect=_fake_ocr):
        chunks = loader.load_and_split_pdf(path, ocr_report=report, force_ocr=True)

    assert {c["page"] for c in chunks} == {1, 2, 3}, "三页都应有内容（第二页回退）"
    assert report.get("ocr_failed") == [], "有文本层可回退，不应记为失败"


def test_渲染失败也不中断整文件(text_pdf):
    """PyMuPDF 渲染整页失败时（如文件损坏），同样回退文本层而不是抛异常。"""
    report = {}
    with patch("loader._render_pdf_page_png", side_effect=RuntimeError("渲染失败")):
        chunks = loader.load_and_split_pdf(text_pdf, ocr_report=report, force_ocr=True)

    joined = "".join(c["text"] for c in chunks)
    assert "Page one has a normal text layer" in joined
    assert report.get("ocr_failed") == []


# ---------------- 默认行为必须与 v3.1 一致（回归保护） ----------------

def test_不勾选强制OCR时文本页不触发OCR(text_pdf):
    """默认行为回归：有正常文本的页一律不走 OCR。"""
    with patch("loader._ocr_image") as mock_ocr:
        report = {}
        chunks = loader.load_and_split_pdf(text_pdf, ocr_report=report)

    assert chunks
    assert not mock_ocr.called, "默认模式不应调用 OCR"
    assert report.get("ocr_pages") == []
    assert report.get("ocr_mode") == "auto"


def test_不勾选强制OCR时图片页仍自动OCR(image_pdf):
    """默认行为回归：无文本页的自动 OCR 兜底不能被破坏。"""
    with patch("loader._ocr_image", return_value="扫描件正文识别结果，长度足够。"):
        report = {}
        chunks = loader.load_and_split_pdf(image_pdf, ocr_report=report)

    assert chunks
    assert report.get("ocr_pages") == [1]


# ---------------- load_and_split_file 的参数透传 ----------------

def test_分发函数能把force_ocr传给PDF(text_pdf):
    """load_and_split_file 应把 force_ocr 透传给 PDF 解析分支。"""
    with patch("loader._ocr_image", return_value="识别文本足够长以通过阈值校验。") as m:
        report = {}
        chunks = loader.load_and_split_file(
            text_pdf, "a.pdf", ocr_report=report, force_ocr=True
        )

    assert m.call_count == 2
    assert report.get("ocr_mode") == "force"
    assert chunks


def test_分发函数对PPTX忽略force_ocr(tmp_path):
    """PPTX 本身有文本结构，force_ocr 对它无效，不应报错也不应触发 OCR。"""
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "PPT标题内容足够长以通过清洗"
    path = str(tmp_path / "d.pptx")
    prs.save(path)

    with patch("loader._ocr_image") as mock_ocr:
        chunks = loader.load_and_split_file(path, "d.pptx", force_ocr=True)

    assert chunks
    assert not mock_ocr.called


# ---------------- 真实 OCR 引擎（默认跳过） ----------------

def test_真实OCR强制模式能识别中文(tmp_path):
    """
    真实 OCR 冒烟用例：默认跳过，设置 RUN_REAL_OCR=1 启用。
    用 PyMuPDF 内置中文字体生成一页含中文的 PDF，强制 OCR 后应能拿到中文片段。
    """
    if os.environ.get("RUN_REAL_OCR") != "1":
        pytest.skip(
            "跳过真实 OCR 用例：默认不加载识别模型。"
            "设置环境变量 RUN_REAL_OCR=1 后可启用（首次会自动下载模型，耗时较长）。"
        )
    try:
        loader._get_ocr()
    except Exception as e:
        pytest.skip(f"跳过：OCR 引擎不可用（{e}）。原因通常是识别模型未下载或依赖未安装。")

    fitz = _open_fitz()
    doc = fitz.open()
    page = doc.new_page(width=800, height=400)
    # china-s 是 PyMuPDF 内置的简体中文字体，避免依赖外部字体文件
    page.insert_text((60, 120), "访问控制列表配置实验结果", fontname="china-s", fontsize=32)
    page.insert_text((60, 220), "基本ACL与高级ACL的区别", fontname="china-s", fontsize=32)
    path = str(tmp_path / "cn.pdf")
    doc.save(path)
    doc.close()

    report = {}
    chunks = loader.load_and_split_pdf(path, ocr_report=report, force_ocr=True)

    assert report.get("ocr_pages") == [1], "强制模式应把该页记入 ocr_pages"
    assert chunks, "应产出可入库的片段"
    assert all(c["page"] == 1 for c in chunks)
    joined = "".join(c["text"] for c in chunks)
    assert "访问控制" in joined or "ACL" in joined, f"识别结果应包含原文关键字，实际：{joined}"
